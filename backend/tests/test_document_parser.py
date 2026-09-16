"""资料读取金丝雀测试：真实文件字节进入解析器，检查正文、顺序和可读错误。

样本在内存生成，既不上传私人攻略，也不依赖Word、网络或收费模型。
如果解析器丢页码、跳过表格、混淆标题或吞掉损坏文件，这些测试必须失败。
"""

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.document.parser import DocumentParseError, parse_document

"""生成两页带文字的PDF或一页无文字PDF；固定文本便于独立核对页码。"""


def pdf_bytes(blank: bool = False) -> bytes:
    writer = PdfWriter()
    for text in [""] if blank else ["West Lake", "Lingyin Temple"]:
        page = writer.add_blank_page(width=300, height=300)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


"""生成标题、段落、表格、二级标题交错的DOCX，检测是否保持阅读顺序。"""


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("杭州", level=1)
    document.add_paragraph("西湖步行")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "交通"
    table.cell(0, 1).text = "地铁"
    document.add_heading("餐饮", level=2)
    document.add_paragraph("面馆")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


"""PDF页码来自原文件，读取结果不能自行合并页或从0开始编号。"""


def test_pdf_keeps_page_numbers() -> None:
    result = parse_document(pdf_bytes(), ".pdf")
    assert [(part.text, part.page_number, part.order) for part in result.sections] == [
        ("West Lake", 1, 1),
        ("Lingyin Temple", 2, 2),
    ]


"""DOCX按正文顺序读表格；没有可靠的排版页码时应返回null而非猜测。"""


def test_docx_keeps_headings_tables_and_order() -> None:
    result = parse_document(docx_bytes(), ".docx")
    assert [part.text for part in result.sections] == [
        "杭州",
        "西湖步行",
        "交通\t地铁",
        "餐饮",
        "面馆",
    ]
    assert result.sections[-1].section_path == ["杭州", "餐饮"]
    assert all(part.page_number is None for part in result.sections)


"""Markdown保留标题层级；标题文本也要保留，不能只剩标题下的正文。"""


def test_markdown_keeps_heading_hierarchy() -> None:
    result = parse_document("# 杭州\n\n西湖\n\n## 交通\n地铁".encode(), ".md")
    assert [part.text for part in result.sections] == ["# 杭州", "西湖", "## 交通", "地铁"]
    assert result.sections[-1].section_path == ["杭州", "交通"]


"""支持常见Windows文本编码；明确的BOM不能被误读成正文乱码。"""


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "gb18030"])
def test_text_encodings(encoding: str) -> None:
    result = parse_document("杭州\r\n\r\n西湖".encode(encoding), ".txt")
    assert [part.text for part in result.sections] == ["杭州", "西湖"]


"""无文字PDF、损坏Office文件、二进制伪装文本和空白文件均不可读成功。"""


@pytest.mark.parametrize(
    ("content", "suffix", "message"),
    [
        (pdf_bytes(blank=True), ".pdf", "扫描"),
        (b"%PDF-broken", ".pdf", "损坏"),
        (b"not a zip", ".docx", "损坏"),
        (b"\x00\x01\x02", ".txt", "文本"),
        (b" \n\t", ".txt", "正文"),
    ],
)
def test_unreadable_files_have_useful_errors(content: bytes, suffix: str, message: str) -> None:
    with pytest.raises(DocumentParseError, match=message):
        parse_document(content, suffix)


"""加密PDF不接受密码，也不能把加密异常原样返回给用户。"""


def test_encrypted_pdf_is_explicitly_rejected() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("private-password")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(DocumentParseError, match="加密"):
        parse_document(output.getvalue(), ".pdf")


"""标题会随每个段落保存，输出预算必须计入这些重复路径，阻止小文件放大成巨型JSON。"""


def test_repeated_long_heading_counts_toward_output_limit() -> None:
    content = ("# " + "题" * 200_000 + "\n\n" + "段\n\n" * 10).encode()
    with pytest.raises(DocumentParseError, match="100万"):
        parse_document(content, ".md")


"""嵌套表格也是正文，不能只读cell.text而静默漏掉单元格内第二层表格。"""


def test_docx_nested_table_content_is_kept() -> None:
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "外层"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "内层景点"
    cell.add_paragraph("尾段")
    output = BytesIO()
    document.save(output)
    sections = parse_document(output.getvalue(), ".docx").sections
    assert sections[0].text == "外层\n内层景点\n尾段"


"""合并单元格只读一次；层层嵌套也不能将同一段文字重复指数放大。"""

def test_docx_nested_merged_cells_are_not_duplicated() -> None:
    document = Document()
    parent = document
    for _ in range(8):
        table = parent.add_table(rows=1, cols=2)
        parent = table.cell(0, 0).merge(table.cell(0, 1))
    parent.text = "唯一景点"
    output = BytesIO()
    document.save(output)
    result = parse_document(output.getvalue(), ".docx")
    assert result.sections[0].text == "唯一景点"
